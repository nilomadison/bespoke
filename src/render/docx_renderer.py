import io

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def render_to_docx(generated: dict, profile: dict) -> DocxDocument:
    """Render a generated resume dict to a python-docx Document.

    ATS rules enforced: no tables, no text boxes, Calibri 11pt body,
    0.75" margins, all content in body paragraphs or bullet lists.
    """
    doc = Document()

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Default font
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)

    # --- Header ---
    name_para = doc.add_paragraph()
    name_run = name_para.add_run(profile.get("full_name", ""))
    name_run.bold = True
    name_run.font.size = Pt(16)
    name_run.font.name = "Calibri"
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    contact_parts = [
        profile.get("email"),
        profile.get("phone"),
        profile.get("location"),
        profile.get("linkedin_url"),
    ]
    contact_line = "  |  ".join(p for p in contact_parts if p)
    if contact_line:
        contact_para = doc.add_paragraph(contact_line)
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_para.runs[0].font.size = Pt(10)

    # --- Summary ---
    if generated.get("summary"):
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(generated["summary"])

    # --- Skills ---
    if generated.get("skills"):
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(generated["skills"])

    # --- Experience ---
    if generated.get("experience"):
        doc.add_heading("Experience", level=1)
        for exp in generated["experience"]:
            job_para = doc.add_paragraph()
            title_run = job_para.add_run(
                f"{exp.get('title', '')}  —  {exp.get('company', '')}"
            )
            title_run.bold = True
            title_run.font.name = "Calibri"
            dates = exp.get("dates", "")
            if dates:
                job_para.add_run(f"   {dates}").font.size = Pt(10)

            for bullet in exp.get("bullets", []):
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(bullet)

    # --- Projects ---
    if generated.get("projects"):
        doc.add_heading("Projects", level=1)
        for proj in generated["projects"]:
            proj_para = doc.add_paragraph()
            name_run = proj_para.add_run(proj.get("name", ""))
            name_run.bold = True
            name_run.font.name = "Calibri"
            desc = proj.get("description", "")
            if desc:
                proj_para.add_run(f": {desc}")

    # --- Education ---
    if generated.get("education"):
        doc.add_heading("Education", level=1)
        for edu in generated["education"]:
            edu_para = doc.add_paragraph()
            edu_run = edu_para.add_run(
                f"{edu.get('degree', '')}  —  {edu.get('institution', '')}"
            )
            edu_run.bold = True
            edu_run.font.name = "Calibri"
            dates = edu.get("dates", "")
            if dates:
                edu_para.add_run(f"   {dates}").font.size = Pt(10)

    return doc


def docx_to_bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
